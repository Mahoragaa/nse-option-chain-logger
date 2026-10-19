from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from nse_oi_logger.orchestrator import OptionChainLogger


def test_cycle_processes_each_strike_expiry_pair(monkeypatch, tmp_path):
    first_expiry = (date.today() + timedelta(days=3)).isoformat()
    second_expiry = (date.today() + timedelta(days=10)).isoformat()
    requested_expiries = []

    class FakeFetcher:
        def fetch_chain(self, symbol, expiry):
            assert symbol == "NIFTY"
            requested_expiries.append(expiry)
            nse_expiry = datetime.strptime(expiry, "%Y-%m-%d").strftime("%d-%b-%Y")
            return {
                "records": {
                    "data": [
                        {
                            "strikePrice": strike,
                            "expiryDate": nse_expiry,
                            "CE": {"openInterest": 100},
                            "PE": {"openInterest": 120},
                        }
                        for strike in (23000, 22900)
                    ]
                }
            }

    monkeypatch.setattr("nse_oi_logger.orchestrator.NseFetcher", FakeFetcher)
    logger = OptionChainLogger(
        "NIFTY",
        None,
        [],
        tmp_path / "options.db",
        tmp_path / "output",
        strike_expiries={
            23000: [first_expiry, second_expiry],
            22900: [first_expiry],
        },
    )
    monkeypatch.setattr(logger, "_market_window_allowed", lambda now=None: now is not None)

    result = logger.process_cycle()

    assert requested_expiries == [first_expiry, second_expiry]
    assert "inserted=3" in result
    workbook = load_workbook(result.split("workbook=")[1])
    try:
        assert f"23000_{first_expiry}" in workbook.sheetnames
        assert f"23000_{second_expiry}" in workbook.sheetnames
        assert f"22900_{first_expiry}" in workbook.sheetnames
    finally:
        workbook.close()
