from openpyxl import load_workbook

from nse_oi_logger.excel_writer import build_summary_metrics, write_daily_workbook


def test_build_summary_metrics():
    rows_by_strike = {
        24500: [
            {"call_oi": 1000, "put_oi": 500, "time_slot": "09:15:00"},
            {"call_oi": 1200, "put_oi": 450, "time_slot": "09:18:00"},
        ],
        24600: [
            {"call_oi": 800, "put_oi": 600, "time_slot": "09:18:00"},
        ],
    }

    summary = build_summary_metrics(rows_by_strike)

    assert summary["total_call_oi"] == 2000
    assert summary["total_put_oi"] == 1050
    assert abs(summary["overall_pcr"] - 0.525) < 1e-9
    assert summary["highest_call_oi_strike"] == 24500
    assert summary["highest_put_oi_strike"] == 24600


def test_write_daily_workbook_keeps_clean_columns_and_filters(tmp_path):
    rows_by_strike = {
        24500: [
            {
                "symbol": "NIFTY",
                "expiry": "2026-10-28",
                "strike": 24500,
                "time_slot": "09:18:00",
                "call_oi": 1200,
                "call_change_oi": 90,
                "put_oi": 900,
                "put_change_oi": -60,
            }
        ]
    }

    file_path = write_daily_workbook(tmp_path, "NIFTY", rows_by_strike)
    workbook = load_workbook(file_path)
    sheet = workbook["24500"]

    assert [cell.value for cell in sheet[1]] == [
        "Time",
        "Strike",
        "Call OI",
        "Call Change in OI",
        "Put OI",
        "Put Change in OI",
    ]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.sheet_view.showGridLines is True
    assert sheet["A2"].border.left.style is not None


def test_workbook_separates_same_strike_by_expiry(tmp_path):
    rows_by_strike = {
        (23000, "2026-10-22"): [
            {
                "strike": 23000,
                "expiry": "2026-10-22",
                "time_slot": "09:18:00",
                "call_oi": 100,
                "call_change_oi": 10,
                "put_oi": 120,
                "put_change_oi": 5,
            }
        ],
        (23000, "2026-10-29"): [
            {
                "strike": 23000,
                "expiry": "2026-10-29",
                "time_slot": "09:18:00",
                "call_oi": 200,
                "call_change_oi": 20,
                "put_oi": 220,
                "put_change_oi": 15,
            }
        ],
    }

    file_path = write_daily_workbook(tmp_path, "NIFTY", rows_by_strike)
    workbook = load_workbook(file_path)

    assert "23000_2026-10-22" in workbook.sheetnames
    assert "23000_2026-10-29" in workbook.sheetnames
    assert workbook["23000_2026-10-22"]["G2"].value == "2026-10-22"
    assert workbook["23000_2026-10-29"]["G2"].value == "2026-10-29"
    workbook.close()
